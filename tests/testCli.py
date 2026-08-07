import cgmesFixtures as fx
import openpyxl
import pytest

from cgmes2excel.cli import main


def testConvertsAnExportAndReportsSuccess(tmp_path):
    fx.sampleExport(tmp_path)
    output = tmp_path / "netz.xlsx"
    assert main([str(tmp_path), "--output", str(output)]) == 0
    assert output.exists()
    assert openpyxl.load_workbook(output).sheetnames == ["Stationen", "NETZELEMENTE"]


def testOutputDefaultsToAWorkbookInTheCurrentDirectory(tmp_path, monkeypatch):
    fx.sampleExport(tmp_path)
    monkeypatch.chdir(tmp_path)
    assert main([str(tmp_path)]) == 0
    assert (tmp_path / "cgmes-export.xlsx").exists()


def testMissingInputExitsWithAnErrorCode(tmp_path, capsys):
    assert main([str(tmp_path / "nope")]) == 2


def testAnEmptyDirectoryExitsWithAnErrorCode(tmp_path):
    empty = tmp_path / "empty"
    empty.mkdir()
    assert main([str(empty)]) == 2


def testSeveralInputPathsCanBeGiven(tmp_path):
    paths = fx.sampleExport(tmp_path)
    output = tmp_path / "netz.xlsx"
    assert main([*[str(p) for p in paths], "--output", str(output)]) == 0


def testEquipmentClassesCanBeAddedFromTheCommandLine(tmp_path):
    body = fx.sampleEquipment() + fx.obj("Breaker", "BR", fx.named("Schalter"), fx.ref("Equipment.EquipmentContainer", "VLA380"))
    fx.writeDocument(tmp_path, "EQ.xml", body, profile="EQ")
    output = tmp_path / "netz.xlsx"
    main([str(tmp_path), "--output", str(output), "--include-class", "Breaker"])
    sheet = openpyxl.load_workbook(output)["NETZELEMENTE"]
    assert any(row[3] == "Schalter" for row in sheet.iter_rows(min_row=2, values_only=True))


def testTheExportedClassSetCanBeReplacedFromTheCommandLine(tmp_path):
    fx.sampleExport(tmp_path)
    output = tmp_path / "netz.xlsx"
    main([str(tmp_path), "--output", str(output), "--only-class", "PowerTransformer"])
    sheet = openpyxl.load_workbook(output)["NETZELEMENTE"]
    types = {row[4] for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert types == {"PowerTransformer"}


def testLogFileIsWrittenWithoutAnsiEscapes(tmp_path):
    fx.sampleExport(tmp_path)
    logFile = tmp_path / "run.log"
    main([str(tmp_path), "--output", str(tmp_path / "n.xlsx"), "--log-file", str(logFile)])
    contents = logFile.read_text(encoding="utf-8")
    assert "Conversion completed" in contents
    assert "\x1b[" not in contents


def testTraceFileCanBeRequested(tmp_path):
    fx.sampleExport(tmp_path)
    trace = tmp_path / "trace.txt"
    main([str(tmp_path), "--output", str(tmp_path / "n.xlsx"), "--trace-file", str(trace)])
    assert "Station Anfang" in trace.read_text(encoding="utf-8")


def testMappingDocumentationCanBePrintedWithoutConverting(tmp_path, capsys):
    assert main(["--print-mapping"]) == 0
    printed = capsys.readouterr().out
    assert "Stationen:" in printed
    assert "NETZELEMENTE:" in printed
    assert "Station Anfang" in printed


def testHelpExits(capsys):
    with pytest.raises(SystemExit) as exit:
        main(["--help"])
    assert exit.value.code == 0


def testVerboseModeEmitsDebugRecords(tmp_path):
    fx.sampleExport(tmp_path)
    logFile = tmp_path / "run.log"
    main([str(tmp_path), "--output", str(tmp_path / "n.xlsx"), "--log-file", str(logFile), "--verbose"])
    assert "DEBUG" in logFile.read_text(encoding="utf-8")


def testMalformedXmlDoesNotCrashTheCommand(tmp_path):
    fx.writeDocument(tmp_path, "EQ.xml", fx.sampleEquipment(), profile="EQ")
    (tmp_path / "broken.xml").write_text("<rdf:RDF><oops>", encoding="utf-8")
    assert main([str(tmp_path), "--output", str(tmp_path / "n.xlsx")]) == 0
