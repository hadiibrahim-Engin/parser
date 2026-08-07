import openpyxl
import pytest

from cgmes2excel.diagnostics import Diagnostics
from cgmes2excel.export.excel import writeWorkbook
from cgmes2excel.export.validation import validateWorkbook
from cgmes2excel.mapping.schema import NETZELEMENTE, NETZELEMENTE_COLUMNS, STATIONEN, STATIONEN_COLUMNS, SchemaViolation
from cgmes2excel.mapping.rows import Row


def stationRow(**values):
    return Row(
        sheet="Stationen",
        values=[values.get(column) for column in STATIONEN_COLUMNS],
        traces={},
        columns=STATIONEN_COLUMNS,
        sourceMrid=values.get("ID-GUID intern-1", ""),
    )


def elementRow(**values):
    return Row(
        sheet="NETZELEMENTE",
        values=[values.get(column) for column in NETZELEMENTE_COLUMNS],
        traces={},
        columns=NETZELEMENTE_COLUMNS,
        sourceMrid=values.get("ID-GUID intern-1", ""),
    )


def writeSample(tmpPath, stations=None, elements=None):
    path = tmpPath / "out.xlsx"
    writeWorkbook(
        path,
        {"Stationen": stations or [], "NETZELEMENTE": elements or []},
        Diagnostics(),
    )
    return path


def loadSheet(path, name):
    return openpyxl.load_workbook(path)[name]


def headerOf(path, name):
    sheet = loadSheet(path, name)
    return [cell.value for cell in sheet[1]]


# --- structure ---------------------------------------------------------------


def testWorkbookContainsExactlyTheTwoContractedSheetsInOrder(tmp_path):
    path = writeSample(tmp_path)
    assert openpyxl.load_workbook(path).sheetnames == ["Stationen", "NETZELEMENTE"]


def testStationenHeaderMatchesTheContract(tmp_path):
    assert headerOf(writeSample(tmp_path), "Stationen") == list(STATIONEN_COLUMNS)


def testNetzelementeHeaderMatchesTheContract(tmp_path):
    assert headerOf(writeSample(tmp_path), "NETZELEMENTE") == list(NETZELEMENTE_COLUMNS)


def testAnEmptyExportStillWritesBothHeaderRows(tmp_path):
    path = writeSample(tmp_path)
    assert loadSheet(path, "Stationen").max_row == 1
    assert loadSheet(path, "NETZELEMENTE").max_row == 1


def testRowsAreWrittenBelowTheHeader(tmp_path):
    path = writeSample(tmp_path, stations=[stationRow(Stationname="Alpha"), stationRow(Stationname="Beta")])
    sheet = loadSheet(path, "Stationen")
    assert sheet.max_row == 3
    assert sheet.cell(row=2, column=3).value == "Alpha"
    assert sheet.cell(row=3, column=3).value == "Beta"


def testColumnCountOnDiskMatchesTheContract(tmp_path):
    path = writeSample(tmp_path, stations=[stationRow(Stationname="Alpha")])
    assert loadSheet(path, "Stationen").max_column == 22
    assert loadSheet(path, "NETZELEMENTE").max_column == 27


# --- value fidelity ----------------------------------------------------------


def testGuidsAreStoredAsTextNotNumbers(tmp_path):
    guid = "0472a783-c766-11e1-8775-005056c00008"
    path = writeSample(tmp_path, stations=[stationRow(**{"ID-GUID intern-1": guid})])
    cell = loadSheet(path, "Stationen").cell(row=2, column=12)
    assert cell.value == guid
    assert cell.data_type == "s"


def testPurelyNumericIdentifiersKeepTheirLeadingZerosAsText(tmp_path):
    path = writeSample(tmp_path, stations=[stationRow(**{"ID-GUID intern-1": "007041"})])
    cell = loadSheet(path, "Stationen").cell(row=2, column=12)
    assert cell.value == "007041"
    assert cell.data_type == "s"
    assert cell.number_format == "@"


def testCoordinatesAreStoredAsNumbers(tmp_path):
    path = writeSample(tmp_path, stations=[stationRow(lat=53.551086, long=9.993682)])
    sheet = loadSheet(path, "Stationen")
    assert sheet.cell(row=2, column=4).value == 53.551086
    assert sheet.cell(row=2, column=5).value == 9.993682
    assert sheet.cell(row=2, column=4).data_type == "n"


def testNetworkElementVoltageIsStoredAsANumber(tmp_path):
    path = writeSample(tmp_path, elements=[elementRow(Spannung=380.0)])
    cell = loadSheet(path, "NETZELEMENTE").cell(row=2, column=6)
    assert cell.value == 380.0
    assert cell.data_type == "n"


def testStationVoltageStaysTextBecauseItMayListSeveralLevels(tmp_path):
    path = writeSample(tmp_path, stations=[stationRow(Spannung="380/110")])
    cell = loadSheet(path, "Stationen").cell(row=2, column=6)
    assert cell.value == "380/110"
    assert cell.data_type == "s"


def testEmptyValuesStayEmptyRatherThanBecomingTheStringNone(tmp_path):
    path = writeSample(tmp_path, stations=[stationRow(Stationname="Alpha")])
    assert loadSheet(path, "Stationen").cell(row=2, column=1).value is None


def testUnicodeAndSpecialCharactersSurviveTheRoundTrip(tmp_path):
    name = "Umspannwerk Süd & Ost — 380/110 kV"
    path = writeSample(tmp_path, stations=[stationRow(Stationname=name)])
    assert loadSheet(path, "Stationen").cell(row=2, column=3).value == name


def testTextThatLooksLikeAFormulaIsNotTreatedAsOne(tmp_path):
    path = writeSample(tmp_path, stations=[stationRow(Kommentar="=SUM(A1:A2)")])
    cell = loadSheet(path, "Stationen").cell(row=2, column=18)
    assert cell.value == "=SUM(A1:A2)"
    assert cell.data_type == "s"


def testTextThatLooksLikeADateIsNotConvertedToADate(tmp_path):
    path = writeSample(tmp_path, stations=[stationRow(IBN="01.02.2024")])
    cell = loadSheet(path, "Stationen").cell(row=2, column=7)
    assert cell.value == "01.02.2024"
    assert cell.data_type == "s"


def testHeaderRowIsFrozenForReadability(tmp_path):
    path = writeSample(tmp_path)
    assert loadSheet(path, "Stationen").freeze_panes == "A2"


# --- guard rails -------------------------------------------------------------


def testWritingARowOfTheWrongWidthIsRefused(tmp_path):
    broken = Row(sheet="Stationen", values=[None] * 21, traces={}, columns=STATIONEN_COLUMNS)
    with pytest.raises(SchemaViolation):
        writeWorkbook(tmp_path / "out.xlsx", {"Stationen": [broken], "NETZELEMENTE": []}, Diagnostics())


def testWritingAnUnknownSheetIsRefused(tmp_path):
    with pytest.raises(SchemaViolation):
        writeWorkbook(tmp_path / "out.xlsx", {"Tabelle1": []}, Diagnostics())


def testEveryContractedSheetMustBeSupplied(tmp_path):
    with pytest.raises(SchemaViolation):
        writeWorkbook(tmp_path / "out.xlsx", {"Stationen": []}, Diagnostics())


# --- validation of the written file ------------------------------------------


def testValidationPassesForAWorkbookWeJustWrote(tmp_path):
    result = validateWorkbook(writeSample(tmp_path, stations=[stationRow(Stationname="Alpha")]))
    assert result.passed
    assert result.violations == []
    assert result.sheetResults == {"Stationen": True, "NETZELEMENTE": True}


def testValidationRejectsARenamedWorksheet(tmp_path):
    path = writeSample(tmp_path)
    workbook = openpyxl.load_workbook(path)
    workbook["Stationen"].title = "Stationen2"
    workbook.save(path)
    result = validateWorkbook(path)
    assert not result.passed
    assert any("Stationen" in violation for violation in result.violations)


def testValidationRejectsARenamedColumn(tmp_path):
    path = writeSample(tmp_path)
    workbook = openpyxl.load_workbook(path)
    workbook["NETZELEMENTE"].cell(row=1, column=5).value = "Elementtyp"
    workbook.save(path)
    result = validateWorkbook(path)
    assert not result.passed
    assert result.sheetResults["NETZELEMENTE"] is False


def testValidationRejectsAReorderedColumn(tmp_path):
    path = writeSample(tmp_path)
    workbook = openpyxl.load_workbook(path)
    sheet = workbook["Stationen"]
    sheet.cell(row=1, column=4).value = "long"
    sheet.cell(row=1, column=5).value = "lat"
    workbook.save(path)
    assert not validateWorkbook(path).passed


def testValidationRejectsAnExtraSheet(tmp_path):
    path = writeSample(tmp_path)
    workbook = openpyxl.load_workbook(path)
    workbook.create_sheet("Extra")
    workbook.save(path)
    result = validateWorkbook(path)
    assert not result.passed
    assert any("Extra" in violation for violation in result.violations)


def testValidationCountsTheDataRowsItFound(tmp_path):
    path = writeSample(tmp_path, stations=[stationRow(Stationname="A"), stationRow(Stationname="B")])
    assert validateWorkbook(path).rowCounts == {"Stationen": 2, "NETZELEMENTE": 0}
