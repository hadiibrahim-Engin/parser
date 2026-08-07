import logging

import cgmesFixtures as fx
import openpyxl
import pytest

from cgmes2excel.cgmes.profiles import Profile
from cgmes2excel.inputs import InputNotFoundError
from cgmes2excel.pipeline import ConversionOptions, convert


def runConversion(tmpPath, inputs=None, **kw):
    options = ConversionOptions(
        inputs=inputs if inputs is not None else [tmpPath],
        output=tmpPath / "out" / "netz.xlsx",
        **kw,
    )
    return convert(options)


def sheetValues(path, name):
    sheet = openpyxl.load_workbook(path)[name]
    header = [cell.value for cell in sheet[1]]
    return [dict(zip(header, [cell.value for cell in row])) for row in sheet.iter_rows(min_row=2)]


# --- happy path --------------------------------------------------------------


def testConversionProducesAValidWorkbook(tmp_path):
    fx.sampleExport(tmp_path)
    result = runConversion(tmp_path)
    assert result.output.exists()
    assert result.validation.passed


def testAllFiveProfilesAreDetected(tmp_path):
    fx.sampleExport(tmp_path)
    result = runConversion(tmp_path)
    assert {doc.profile for doc in result.documents} == {
        Profile.EQ,
        Profile.SSH,
        Profile.TP,
        Profile.SV,
        Profile.GL,
    }


def testProfilesAreDetectedFromModelHeadersNotFilenames(tmp_path):
    paths = fx.sampleExport(tmp_path)
    for path in paths:
        path.rename(path.with_name(path.name.replace("grid_", "profile_").replace("EQ", "XX")))
    result = runConversion(tmp_path)
    assert Profile.EQ in {doc.profile for doc in result.documents}


def testStationsAreExportedForEverySubstation(tmp_path):
    fx.sampleExport(tmp_path)
    result = runConversion(tmp_path)
    rows = sheetValues(result.output, "Stationen")
    assert {row["Stationname"] for row in rows} == {"Umspannwerk Alpha", "Umspannwerk Beta"}


def testNetworkElementsAreExportedForBranchEquipment(tmp_path):
    fx.sampleExport(tmp_path)
    result = runConversion(tmp_path)
    rows = sheetValues(result.output, "NETZELEMENTE")
    assert {row["Title"] for row in rows} == {"Leitung Alpha-Beta", "Transformator A"}


def testEndToEndStationDerivationCrossesProfiles(tmp_path):
    fx.sampleExport(tmp_path)
    result = runConversion(tmp_path)
    line = next(row for row in sheetValues(result.output, "NETZELEMENTE") if row["Title"] == "Leitung Alpha-Beta")
    assert line["Station Anfang"] == "Umspannwerk Alpha"
    assert line["Station Ende"] == "Umspannwerk Beta"
    assert line["Y-Knoten-1"] == "Knoten A 380"


def testGeographyFromTheGlProfileReachesTheWorkbook(tmp_path):
    fx.sampleExport(tmp_path)
    result = runConversion(tmp_path)
    alpha = next(row for row in sheetValues(result.output, "Stationen") if row["Stationname"] == "Umspannwerk Alpha")
    assert alpha["lat"] == 53.551086
    assert alpha["long"] == 9.993682


def testInputFileOrderDoesNotChangeTheResult(tmp_path):
    paths = fx.sampleExport(tmp_path)
    forward = runConversion(tmp_path, inputs=paths)
    backward = runConversion(tmp_path, inputs=list(reversed(paths)))
    assert sheetValues(forward.output, "NETZELEMENTE") == sheetValues(backward.output, "NETZELEMENTE")


def testRunningTwiceProducesIdenticalRows(tmp_path):
    fx.sampleExport(tmp_path)
    first = sheetValues(runConversion(tmp_path).output, "Stationen")
    second = sheetValues(runConversion(tmp_path).output, "Stationen")
    assert first == second


def testZippedExportsAreConverted(tmp_path):
    import zipfile

    archive = tmp_path / "export.zip"
    with zipfile.ZipFile(archive, "w") as zf:
        zf.writestr("EQ.xml", fx.rdfDocument(fx.sampleEquipment(), profile="EQ"))
        zf.writestr("TP.xml", fx.rdfDocument(fx.sampleTopology(), profile="TP"))
    result = runConversion(tmp_path, inputs=[archive])
    assert result.validation.passed
    assert len(sheetValues(result.output, "Stationen")) == 2


# --- optional profiles -------------------------------------------------------


def testConversionWorksWithoutTheGlProfile(tmp_path):
    fx.writeDocument(tmp_path, "EQ.xml", fx.sampleEquipment(), profile="EQ")
    fx.writeDocument(tmp_path, "TP.xml", fx.sampleTopology(), profile="TP")
    result = runConversion(tmp_path)
    assert result.validation.passed
    alpha = next(row for row in sheetValues(result.output, "Stationen") if row["Stationname"] == "Umspannwerk Alpha")
    assert alpha["lat"] is None


def testConversionWorksWithTheEquipmentProfileAlone(tmp_path):
    fx.writeDocument(tmp_path, "EQ.xml", fx.sampleEquipment(), profile="EQ")
    result = runConversion(tmp_path)
    assert result.validation.passed
    assert len(sheetValues(result.output, "NETZELEMENTE")) == 2


def testMissingProfilesAreVisibleInTheSummary(tmp_path):
    fx.writeDocument(tmp_path, "EQ.xml", fx.sampleEquipment(), profile="EQ")
    result = runConversion(tmp_path)
    assert result.summary.missingProfiles == ["SSH", "TP", "SV", "GL"]


# --- configuration -----------------------------------------------------------


def testExtraEquipmentClassesCanBeIncluded(tmp_path):
    body = fx.sampleEquipment() + fx.obj("Breaker", "BR", fx.named("Schalter"), fx.ref("Equipment.EquipmentContainer", "VLA380"))
    fx.writeDocument(tmp_path, "EQ.xml", body, profile="EQ")
    result = runConversion(tmp_path, includeClasses=["Breaker"])
    assert any(row["Title"] == "Schalter" for row in sheetValues(result.output, "NETZELEMENTE"))


def testTheExportedClassSetCanBeReplaced(tmp_path):
    fx.sampleExport(tmp_path)
    result = runConversion(tmp_path, onlyClasses=["PowerTransformer"])
    rows = sheetValues(result.output, "NETZELEMENTE")
    assert {row["Element Typ"] for row in rows} == {"PowerTransformer"}


# --- diagnostics -------------------------------------------------------------


def testSummaryCountsParsedObjectsPerClass(tmp_path):
    fx.sampleExport(tmp_path)
    result = runConversion(tmp_path)
    assert result.summary.classCounts["Substation"] == 2
    assert result.summary.classCounts["ACLineSegment"] == 1


def testSummaryCountsOutputRows(tmp_path):
    fx.sampleExport(tmp_path)
    summary = runConversion(tmp_path).summary
    assert summary.stationRows == 2
    assert summary.elementRows == 2


def testSummaryCountsEmptyTargetFields(tmp_path):
    fx.sampleExport(tmp_path)
    assert runConversion(tmp_path).summary.emptyFields > 0


def testUnresolvedReferencesAreCountedAndReported(tmp_path):
    body = fx.sampleEquipment() + fx.obj("Terminal", "TX", fx.ref("Terminal.ConductingEquipment", "GHOST"))
    fx.writeDocument(tmp_path, "EQ.xml", body, profile="EQ")
    result = runConversion(tmp_path)
    assert result.summary.unresolvedReferences >= 1


def testMalformedDocumentIsReportedWithoutAbortingTheRun(tmp_path):
    fx.writeDocument(tmp_path, "EQ.xml", fx.sampleEquipment(), profile="EQ")
    (tmp_path / "broken.xml").write_text("<rdf:RDF><oops>", encoding="utf-8")
    result = runConversion(tmp_path)
    assert result.validation.passed
    assert "malformedDocument" in result.diagnostics.countByCode()


def testAnInputWithoutAnyDocumentIsRejected(tmp_path):
    with pytest.raises(InputNotFoundError):
        runConversion(tmp_path, inputs=[tmp_path / "missing"])


def testAnEmptyInputDirectoryIsReportedClearly(tmp_path):
    from cgmes2excel.pipeline import NoDocumentsError

    empty = tmp_path / "empty"
    empty.mkdir()
    with pytest.raises(NoDocumentsError):
        runConversion(tmp_path, inputs=[empty])


# --- logging -----------------------------------------------------------------


def testProgressIsLoggedAtInfoLevel(tmp_path, caplog):
    fx.sampleExport(tmp_path)
    with caplog.at_level(logging.INFO):
        runConversion(tmp_path)
    assert "Detected profile" in caplog.text
    assert "Stationen" in caplog.text


def testSummaryIsLoggedAtTheEndOfTheRun(tmp_path, caplog):
    fx.sampleExport(tmp_path)
    with caplog.at_level(logging.INFO):
        result = runConversion(tmp_path)
    assert "Conversion completed" in caplog.text
    assert str(result.output) in caplog.text


def testTraceReportCanBeWritten(tmp_path):
    fx.sampleExport(tmp_path)
    tracePath = tmp_path / "trace.txt"
    runConversion(tmp_path, traceFile=tracePath)
    contents = tracePath.read_text(encoding="utf-8")
    assert "Station Anfang" in contents
    assert "TopologicalNode" in contents
