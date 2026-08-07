"""Validation of the workbook that was actually written to disk.

The contract is only met if the file on disk meets it, so validation reopens
the saved workbook rather than trusting the in-memory rows.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from ..mapping.schema import SHEETS


@dataclass(slots=True)
class ValidationResult:
    """Outcome of checking a workbook against the contract."""

    violations: list[str] = field(default_factory=list)
    sheetResults: dict[str, bool] = field(default_factory=dict)
    rowCounts: dict[str, int] = field(default_factory=dict)

    @property
    def passed(self) -> bool:
        return not self.violations


def validateWorkbook(path: Path) -> ValidationResult:
    """Check worksheet names, column names and column order of a saved workbook."""
    result = ValidationResult()
    workbook = openpyxl.load_workbook(path, read_only=True)
    try:
        expected = [schema.name for schema in SHEETS]
        if workbook.sheetnames != expected:
            result.violations.append(
                f"Workbook must contain exactly {expected} in that order, found {workbook.sheetnames}"
            )

        for schema in SHEETS:
            if schema.name not in workbook.sheetnames:
                result.violations.append(f"{schema.name}: worksheet is missing")
                result.sheetResults[schema.name] = False
                result.rowCounts[schema.name] = 0
                continue

            sheet = workbook[schema.name]
            # max_row is unreliable on read-only sheets, so rows are counted.
            rows = sheet.iter_rows(values_only=True)
            header = [value for value in next(rows, ())]
            violations = schema.headerViolations(header)
            result.violations.extend(violations)
            result.sheetResults[schema.name] = not violations
            result.rowCounts[schema.name] = sum(1 for _ in rows)
    finally:
        workbook.close()
    return result
