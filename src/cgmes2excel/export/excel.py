"""Workbook generation.

Formatting decisions must never change a value: text columns are written as
strings with an explicit text format and a pinned string cell type, so
identifiers, leading zeros, date-like strings and values starting with ``=``
cannot be reinterpreted as numbers, dates or formulas.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from cgmes2excel.diagnostics import Diagnostics
from cgmes2excel.logging import getLogger
from cgmes2excel.mapping.rows import Row
from cgmes2excel.mapping.schema import SHEETS, SchemaViolation, SheetSchema, ValueType

_TEXT_FORMAT = "@"
_HEADER_FONT = Font(bold=True)

logger = getLogger(__name__)


def writeWorkbook(path: Path, rowsBySheet: Mapping[str, Sequence[Row]], diagnostics: Diagnostics) -> Path:
    """Write the two contracted worksheets to ``path``."""
    _checkSheets(rowsBySheet)

    workbook = Workbook()
    workbook.remove(workbook.active)
    for schema in SHEETS:
        sheet = workbook.create_sheet(title=schema.name)
        _writeSheet(sheet, schema, rowsBySheet[schema.name])
        logger.info("Wrote %d row(s) to worksheet %s", len(rowsBySheet[schema.name]), schema.name)

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    diagnostics.info("workbookWritten", "Workbook written", path=str(path))
    return path


def _checkSheets(rowsBySheet: Mapping[str, Sequence[Row]]) -> None:
    expected = {schema.name for schema in SHEETS}
    unexpected = sorted(set(rowsBySheet) - expected)
    if unexpected:
        raise SchemaViolation(f"Unknown worksheet(s) requested: {', '.join(unexpected)}")
    missing = sorted(expected - set(rowsBySheet))
    if missing:
        raise SchemaViolation(f"Missing worksheet(s): {', '.join(missing)}")


def _writeSheet(sheet: Worksheet, schema: SheetSchema, rows: Sequence[Row]) -> None:
    for index, column in enumerate(schema.columns, start=1):
        cell = sheet.cell(row=1, column=index, value=column)
        cell.font = _HEADER_FONT
        cell.number_format = _TEXT_FORMAT
    sheet.freeze_panes = "A2"

    for rowNumber, row in enumerate(rows, start=2):
        schema.validateRowWidth(len(row.values))
        for index, (column, value) in enumerate(zip(schema.columns, row.values), start=1):
            _writeCell(sheet, schema, rowNumber, index, column, value)


def _writeCell(
    sheet: Worksheet,
    schema: SheetSchema,
    rowNumber: int,
    columnNumber: int,
    column: str,
    value: object | None,
) -> None:
    if value is None:
        return
    cell = sheet.cell(row=rowNumber, column=columnNumber)
    if schema.valueTypeOf(column) is ValueType.NUMBER:
        cell.value = value
        return
    cell.number_format = _TEXT_FORMAT
    cell.value = value if isinstance(value, str) else str(value)
    # Pinned after assignment: openpyxl would otherwise type a leading '=' as a formula.
    cell.data_type = "s"
